"""One slide per traced neuron: reconstruction, probability at three thresholds, image.

Usage: uv run --with python-pptx python scripts/deck_cells.py <figure dir> <recall npz> <out.pptx>
"""

import sys
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

fig_dir, recall_npz, out = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
z = np.load(recall_npz)
val, own, names = z["val"], z["owner"], list(z["names"])
INK, MUTED, PAGE = (
    RGBColor(0x20, 0x20, 0x20),
    RGBColor(0x6B, 0x66, 0x63),
    RGBColor(0xF3, 0xF0, 0xE8),
)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, s, size, color=INK, font="Allen Institute Headline", bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h).text_frame
    tb.word_wrap = True
    p = tb.paragraphs[0]
    p.text = s
    p.font.size, p.font.name, p.font.bold, p.font.color.rgb = Pt(size), font, bold, color


# title slide
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAGE
text(
    s,
    Inches(0.6),
    Inches(2.4),
    Inches(12),
    Inches(1.2),
    "794495 · u-net probability along 22 traced neurons/",
    36,
)
overall = {t: np.mean(val >= round(t * 255)) for t in (0.2, 0.5, 0.9)}
text(
    s,
    Inches(0.6),
    Inches(3.6),
    Inches(12),
    Inches(2.4),
    "Recall of 6.8 m of traced axon: "
    + " · ".join(f"p >= {t}: {r:.1%}" for t, r in overall.items())
    + "\nEach slide: the reconstruction, the u-net probability (z max projection over the "
    "132-voxel chunks the reconstruction passes through) at p >= 0.2, 0.5 and 0.9, and the image "
    "over the same chunks.",
    16,
    MUTED,
    "Allen Institute Text",
)
for k, name in enumerate(names):
    f = fig_dir / f"{name}.png"
    if not f.exists():
        continue
    s = prs.slides.add_slide(blank)
    m = own == k
    rec = " · ".join(f"p >= {t}: {np.mean(val[m] >= round(t * 255)):.0%}" for t in (0.2, 0.5, 0.9))
    text(s, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.6), f"794495 · {name}/", 26)
    text(
        s,
        Inches(0.4),
        Inches(0.85),
        Inches(12.5),
        Inches(0.4),
        f"traced axon recalled: {rec}",
        14,
        MUTED,
        "Allen Institute Text",
    )
    w, h = Image.open(f).size
    maxw, maxh = prs.slide_width - Inches(0.6), prs.slide_height - Inches(1.5)
    scale = min(maxw / w, maxh / h)
    pw, ph = int(w * scale), int(h * scale)
    s.shapes.add_picture(str(f), Emu((prs.slide_width - pw) // 2), Inches(1.35), Emu(pw), Emu(ph))
prs.save(out)
print("wrote", out, len(prs.slides), "slides")

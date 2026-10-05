"""Final deck: every sample, before and after removing predicted cable under manual traces.

Usage: uv run --with python-pptx python scripts/deck_before_after.py <batch dir> <out.pptx>
Needs render_untraced.py output (qc/<tag>_before_after.png) and untraced_summary.json.
"""

from __future__ import annotations

import datetime
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

INK, SLATE, MUTED = "000000", "666666", "AAA39F"  # Allen ink / muted / warm gray
PAPER, TINT, LIGHT, BORDER = "FFFFFF", "F3F0E8", "E6E6FF", "DED9D1"
BLUE, VIOLET, AXON = "6464FF", "8246E1", "FF6E00"
HEAD, BODY = "Allen Institute Headline", "Allen Institute Text"  # no µ / → glyphs

batch, out_path = Path(sys.argv[1]), Path(sys.argv[2])
jobs = {j["subject"]: j for j in json.loads((batch / "jobs.json").read_text())}
report = json.loads((batch / "untraced_summary.json").read_text())
img_dir = batch / "deck" / "img_before_after"
img_dir.mkdir(parents=True, exist_ok=True)

samples, failed = [], []
for subject in sorted(jobs):
    src = batch / subject / "qc" / f"{subject}_MD_L2_before_after.png"
    if not src.exists():
        failed.append(subject)
        continue
    im = Image.open(src).convert("RGB")
    im.thumbnail((2600, 2600))
    dst = img_dir / f"{subject}.jpg"
    im.save(dst, quality=85)
    samples.append({"subject": subject, "image": dst, "w": im.width, "h": im.height,
                    "stats": report.get(subject)})  # fmt: skip
removed = [s for s in samples if s["stats"] and s["stats"]["removed_um"] > 0]
kept = [s for s in samples if s["stats"] and s["stats"]["removed_um"] <= 0]
untraced = [s for s in samples if not s["stats"]]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def background(slide, colour: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(colour)


def text(slide, x, y, w, h, value, *, size, colour, font=BODY, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(value if isinstance(value, list) else [value]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
        run.font.color.rgb = rgb(colour)
    return box


def card(slide, x, y, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w),
                               Inches(h))  # fmt: skip
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(TINT)
    s.line.color.rgb = rgb(BORDER)
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    s.adjustments[0] = 0.05


# title
s = prs.slides.add_slide(BLANK)
background(s, BLUE)
text(s, 0.8, 2.0, 11.9, 1.0, "md skeletons, traced neurites removed/", size=34,
     colour=PAPER, font=HEAD, bold=True)  # fmt: skip
text(s, 0.8, 3.1, 11.7, 0.9, "exaSPIM · fine-tuned TRAILMAP U-Net · predicted cable within 10 um "
     "of a manual (Horta) trace removed, with the stubs it leaves (< 100 um)", size=20,
     colour=LIGHT)  # fmt: skip
total_removed = sum(x["stats"]["removed_um"] for x in removed) / 1e3
for i, (value, label) in enumerate((
    (str(len(samples)), "samples"),
    (str(len(removed) + len(kept)), "with manual traces"),
    (f"{total_removed:,.0f} mm", "of traced cable removed"),
)):  # fmt: skip
    text(s, 0.8 + i * 3.9, 4.5, 3.6, 0.9, value, size=44, colour=PAPER, font=HEAD, bold=True)
    text(s, 0.8 + i * 3.9, 5.45, 3.6, 0.4, label, size=16, colour=LIGHT)
text(s, 0.8, 6.7, 4, 0.3, datetime.date.today().isoformat(), size=12, colour=LIGHT)

# summary of groups
s = prs.slides.add_slide(BLANK)
background(s, PAPER)
text(s, 0.6, 0.45, 12, 0.8, "which samples had traces/", size=36, colour=INK, font=HEAD, bold=True)
groups = (
    (f"traces removed · {len(removed)}", removed, VIOLET),
    (f"traces outside MD · {len(kept)}", kept, SLATE),
    (f"no manual traces · {len(untraced)}", untraced, MUTED),
)
for i, (head, members, colour) in enumerate(groups):
    x = 0.6 + i * 4.1
    card(s, x, 1.5, 3.85, 5.4)
    text(s, x + 0.25, 1.7, 3.4, 0.5, head, size=18, colour=colour, font=HEAD, bold=True)
    ids = [m["subject"] for m in members]
    cols = [ids[k::3] for k in range(3)] if len(ids) > 12 else [ids]
    for c, col in enumerate(cols):
        text(s, x + 0.25 + c * 1.15, 2.3, 1.1, 4.4, col, size=13, colour=INK)

# one slide per sample
for x in samples:
    s = prs.slides.add_slide(BLANK)
    background(s, PAPER)
    text(s, 0.5, 0.3, 4, 0.7, x["subject"] + "/", size=32, colour=INK, font=HEAD, bold=True)
    st = x["stats"]
    if st is None:
        status, colour = "no manual traces · nothing removed", MUTED
    elif st["removed_um"] <= 0:
        status, colour = (
            f"{st['trace_files']} manual trace files, none in MD · nothing removed",
            SLATE,
        )
    else:
        status = (f"{st['trace_files']} manual trace files · {st['removed_um'] / 1e3:,.0f} mm of "
                  f"predicted cable removed · axon candidates "
                  f"{st['before']['axon_candidate_um'] / 1e3:,.0f} -> "
                  f"{st['after']['axon_candidate_um'] / 1e3:,.0f} mm")  # fmt: skip
        colour = VIOLET
    text(s, 0.5, 0.95, 12.3, 0.35, status, size=14, colour=colour)
    w = 12.3
    h = w * x["h"] / x["w"]
    if h > 5.9:
        h, w = 5.9, 5.9 * x["w"] / x["h"]
    s.shapes.add_picture(str(x["image"]), Inches((13.333 - w) / 2), Inches(1.4), Inches(w),
                         Inches(h))  # fmt: skip
    s.notes_slide.notes_text_frame.text = (
        f"{x['subject']}: {status}. Asset {jobs[x['subject']]['asset']}"
    )

# closing
s = prs.slides.add_slide(BLANK)
background(s, BLUE)
text(s, 0.8, 0.8, 11.7, 0.8, "how to read the figures/", size=36, colour=PAPER, font=HEAD,
     bold=True)  # fmt: skip
notes = [
    "• Left (before): predicted skeleton (orange = axon candidates, pale blue = dendrites of "
    "detected somata), manual traces in violet, removed cable in maroon",
    "• Right (after): what is left once cable within 10 um of a trace and the < 100 um stubs it "
    "leaves are removed",
    "• Beige = MD (CCF 362) in sample space; dashed = MD grown by 20% of its volume",
    *([f"• Not shown: {', '.join(failed)} (segmentation run failed)"] if failed else []),
]  # fmt: skip
box = text(s, 0.8, 1.9, 11.7, 4.5, notes, size=18, colour=LIGHT)
for p in box.text_frame.paragraphs:
    p.space_after = Pt(12)

prs.save(out_path)
print("wrote", out_path, len(prs.slides), "slides;", len(removed), "removed,", len(kept),
      "traced not in MD,", len(untraced), "untraced,", len(failed), "missing")  # fmt: skip

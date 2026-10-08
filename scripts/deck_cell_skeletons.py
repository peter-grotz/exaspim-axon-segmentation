"""One slide per traced neuron: ground truth vs the full-volume skeleton.

Usage:
    uv run --with python-pptx python scripts/deck_cell_skeletons.py <fig dir> <gt_eval.npz> <out>
"""

import sys
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

fig_dir, eval_npz, out = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
z = np.load(eval_npz)
dg, w, own, names = z["dg"], z["W"], z["O"], list(z["names"])
INK, MUTED = RGBColor(0x20, 0x20, 0x20), RGBColor(0x6B, 0x66, 0x63)
PAGE = RGBColor(0xF3, 0xF0, 0xE8)


def recall(mask, tol):
    return float(np.sum(w[mask] * (dg[mask] <= tol)) / w[mask].sum())


def text(slide, x, y, width, height, s, size, color=INK, font="Allen Institute Headline"):
    frame = slide.shapes.add_textbox(x, y, width, height).text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = s
    p.font.size, p.font.name, p.font.color.rgb = Pt(size), font, color


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAGE
text(s, Inches(0.6), Inches(2.4), Inches(12), Inches(1.2),
     "794495 · full-volume skeleton vs 22 traced neurons/", 36)  # fmt: skip
every = np.ones(len(own), bool)
summary = " · ".join(f"within {t} um {recall(every, t):.0%}" for t in (3, 5, 10))
text(s, Inches(0.6), Inches(3.6), Inches(12), Inches(2.4),
     f"Recall of {w.sum() / 1e6:.2f} m of traced axon: {summary}.\nEach slide: the reconstruction "
     "(axon recalled within 5 um orange, missed red, dendrite grey), the predicted skeleton within "
     "10 um of it, and the overlay.", 16, MUTED, "Allen Institute Text")  # fmt: skip
for k, name in enumerate(names):
    f = fig_dir / f"{name}.png"
    if not f.exists():
        continue
    s = prs.slides.add_slide(blank)
    m = own == k
    rec = " · ".join(f"{t} um: {recall(m, t):.0%}" for t in (3, 5, 10))
    text(s, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.6), f"794495 · {name}/", 26)
    line = f"traced axon recalled by the skeleton within {rec}"
    font = "Allen Institute Text"
    text(s, Inches(0.4), Inches(0.85), Inches(12.5), Inches(0.4), line, 14, MUTED, font)
    wpx, hpx = Image.open(f).size
    maxw, maxh = prs.slide_width - Inches(0.6), prs.slide_height - Inches(1.5)
    scale = min(maxw / wpx, maxh / hpx)
    pw, ph = int(wpx * scale), int(hpx * scale)
    s.shapes.add_picture(str(f), Emu((prs.slide_width - pw) // 2), Inches(1.35), Emu(pw), Emu(ph))
prs.save(out)
print("wrote", out, len(prs.slides), "slides")

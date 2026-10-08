"""One slide per traced neuron: the skeleton fragments touching it, splits and merges.

Usage:
    uv run --with python-pptx python scripts/deck_cell_fragments.py <fig dir> <gt_eval.npz> <out>
"""

import json
import sys
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

fig_dir, eval_npz, out = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
z = np.load(eval_npz)
dg, w, own, names = z["dg"], z["W"], z["O"], list(z["names"])
INK, MUTED = RGBColor(0x20, 0x20, 0x20), RGBColor(0x6B, 0x66, 0x63)
PAGE = RGBColor(0xF3, 0xF0, 0xE8)


def recall(mask, tol):
    return float(np.sum(w[mask] * (dg[mask] <= tol)) / w[mask].sum())


def text(slide, x, y, width, height, s, size, color=INK, font="Allen Institute Headline"):
    frame = slide.shapes.add_textbox(x, y, width, height).text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = s
    p.font.size, p.font.name, p.font.color.rgb = Pt(size), font, color


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAGE
text(s, Inches(0.6), Inches(2.4), Inches(12), Inches(1.2),
     "794495 · skeleton fragments on the 22 traced neurons: splits and merges/", 32)  # fmt: skip
stats_file = Path(eval_npz).parent / "fragment_stats.json"
stats = {r["cell"]: r for r in json.loads(stats_file.read_text())}
intro = "\n".join(
    [
        "U-Net + block-wise skeleton only (no tracer, default cleanup). A fragment is one",
        "connected skeleton piece; it touches a cell where an edge is within 10 um of the trace.",
        "Left: the cell's own cable coloured by fragment (splits).",
        "Right: all cable in those fragments: on the cell (blue), on other traced cells (violet),",
        "untraced (red), i.e. merge errors.",
        "Largest components: 14.4 m, 9.8 m and 7.6 m (the last joins 12 of the traced cells).",
    ]
)
text(s, Inches(0.6), Inches(3.5), Inches(12), Inches(3), intro, 16, MUTED, "Allen Institute Text")
for k, name in enumerate(names):
    f = fig_dir / f"{name}.png"
    if not f.exists():
        continue
    s = prs.slides.add_slide(blank)
    m = own == k
    st = stats[name]
    rec = (
        f"recall 5 um {recall(m, 5):.0%} · {st['fragments_100um']} fragments >= 100 um · "
        f"largest fragment {st['largest_fragment_m']} m · joined to "
        f"{st['cells_in_its_fragments']} other traced cells"
    )
    text(s, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.6), f"794495 · {name}/", 26)
    line = rec
    font = "Allen Institute Text"
    text(s, Inches(0.4), Inches(0.85), Inches(12.5), Inches(0.4), line, 14, MUTED, font)
    wpx, hpx = Image.open(f).size
    maxw, maxh = prs.slide_width - Inches(0.6), prs.slide_height - Inches(1.5)
    scale = min(maxw / wpx, maxh / hpx)
    pw, ph = int(wpx * scale), int(hpx * scale)
    s.shapes.add_picture(str(f), Emu((prs.slide_width - pw) // 2), Inches(1.35), Emu(pw), Emu(ph))
prs.save(out)
print("wrote", out, len(prs.slides), "slides")

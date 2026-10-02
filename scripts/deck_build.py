"""Build the batch deck from <batch>/deck/samples.json (see scripts/deck_prepare.py).

Usage: uv run --with python-pptx python scripts/deck_build.py <samples.json> <out.pptx>
"""

from __future__ import annotations

import datetime
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK, SLATE, MUTED = "1B2430", "5B6675", "8A93A0"
PAPER, TINT, LIGHT = "FFFFFF", "F3F5F8", "C9D2DD"
AXON, DENDRITE, DENDRITE_TEXT = "FF5400", "9EBCDB", "4F7FAE"
HEAD, BODY = "Cambria", "Calibri"

samples = json.loads(Path(sys.argv[1]).read_text())
ok = [s for s in samples if s["ok"]]
failed = [s for s in samples if not s["ok"]]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def background(slide, colour: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(colour)


def text(slide, x, y, w, h, value, *, size, colour, font=BODY, bold=False, align=None,
         anchor=MSO_ANCHOR.TOP):  # fmt: skip
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        if align:
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
        run.font.color.rgb = rgb(colour)
    return box


def card(slide, x, y, w, h, colour=TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(colour)
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def fmt(x, digits=0):
    return "–" if x is None else f"{x:,.{digits}f}"


def total(key):
    return sum(s.get(key) or 0 for s in ok)


# 1 · title
s = prs.slides.add_slide(BLANK)
background(s, INK)
text(s, 0.8, 2.0, 11.7, 1.0, "TRAILMAP axon segmentation across exaSPIM", size=40, colour=PAPER,
     font=HEAD, bold=True)  # fmt: skip
text(s, 0.8, 3.05, 11.7, 0.6, "Mediodorsal thalamus (MD) · fine-tuned U-Net · skeletons, debris "
     "removal, dendrite labelling", size=20, colour=LIGHT)  # fmt: skip
for i, (value, label, colour) in enumerate((
    (fmt(len(ok)), "samples processed", PAPER),
    (f"{total('skeleton_mm') / 1000:,.1f} m", "of skeleton in MD", PAPER),
    (f"{fmt(total('axon_mm'))} mm", "axon candidates", AXON),
)):  # fmt: skip
    text(s, 0.8 + i * 3.9, 4.5, 3.6, 0.9, value, size=44, colour=colour, font=HEAD, bold=True)
    text(s, 0.8 + i * 3.9, 5.45, 3.6, 0.4, label, size=16, colour=LIGHT)
text(s, 0.8, 6.7, 4, 0.3, datetime.date.today().isoformat(), size=12, colour=MUTED)

# 2 · pipeline
s = prs.slides.add_slide(BLANK)
background(s, PAPER)
text(s, 0.6, 0.45, 12, 0.8, "What every sample went through", size=36, colour=INK, font=HEAD,
     bold=True)  # fmt: skip
steps = (
    ("Region", "MD located in the reverse-registered CCF annotation and grown by 20% of its "
     "volume (~58 µm shell)."),
    ("Segment", "Fine-tuned TRAILMAP U-Net (v2) on fused level 2 (4 × 3 × 3 µm), on a Code Ocean "
     "GPU."),
    ("Skeletonize", "p > 0.2, thinned; speckle, short and isolated pieces removed as debris."),
    ("Connect", "Loose ends joined to aligned cable ≤ 50 µm away along the brightest path."),
    ("Label", "Cable connected to, or within the field of, a detected soma → dendrite; the rest "
     "→ axon candidates."),
)  # fmt: skip
for i, (head, body) in enumerate(steps):
    x = 0.6 + i * 2.48
    card(s, x, 1.75, 2.28, 3.9)
    card(s, x + 0.25, 2.0, 0.62, 0.62, AXON if i == 4 else INK, MSO_SHAPE.OVAL)
    text(s, x + 0.25, 2.0, 0.62, 0.62, str(i + 1), size=20, colour=PAPER, font=HEAD, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    text(s, x + 0.25, 2.85, 1.85, 0.45, head, size=20, colour=INK, font=HEAD, bold=True)
    text(s, x + 0.25, 3.35, 1.85, 2.1, body, size=14, colour=SLATE)
text(s, 0.6, 6.05, 12.1, 0.6, "Outputs per sample: one Horta-importable SWC per piece (type 3 = "
     "dendrite, 2 = axon candidate), axon candidates separately, QC figures.", size=14,
     colour=SLATE)  # fmt: skip

# 3 · summary chart
s = prs.slides.add_slide(BLANK)
background(s, PAPER)
text(s, 0.6, 0.45, 12, 0.8, "Cable found in MD, per sample", size=36, colour=INK, font=HEAD,
     bold=True)  # fmt: skip
text(s, 0.6, 1.2, 12.1, 0.4, "Skeleton length inside the grown MD (mm, log scale): dendrites of "
     "detected somata and axon candidates, samples ordered by total length.", size=14,
     colour=SLATE)  # fmt: skip
ordered = sorted(ok, key=lambda x: -x["skeleton_mm"])
data = CategoryChartData()
data.categories = [x["subject"] for x in ordered]
# log axis: lengths span 1 mm to several metres; empty categories get no bar (log 0 is undefined)
data.add_series(
    "Dendrites of detected somata",
    [x["dendrite_mm"] if (x["dendrite_mm"] or 0) >= 1 else None for x in ordered],
)
data.add_series(
    "Axon candidates", [x["axon_mm"] if (x["axon_mm"] or 0) >= 1 else None for x in ordered]
)
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.7), Inches(12.1),
                           Inches(5.5), data).chart  # fmt: skip
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.TOP
chart.legend.include_in_layout = False
chart.legend.font.size = Pt(12)
for series, colour in zip(chart.series, (DENDRITE, AXON), strict=True):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(colour)
chart.plots[0].gap_width = 30
chart.plots[0].overlap = 0
cat, val = chart.category_axis, chart.value_axis
cat.tick_labels.font.size = Pt(8)
# every sample label, written vertically
skip = cat._element.makeelement(qn("c:tickLblSkip"), {"val": "1"})
after = cat._element.find(qn("c:noMultiLvlLbl"))  # schema order: tickLblSkip precedes it
after.addprevious(skip) if after is not None else cat._element.append(skip)
tx = cat._element.find(qn("c:txPr"))
tx.find(qn("a:bodyPr")).set("rot", "-5400000")
cat.tick_labels.font.color.rgb = rgb(SLATE)
cat.format.line.color.rgb = rgb("C9D2DD")
val.tick_labels.font.size = Pt(10)
val.tick_labels.font.color.rgb = rgb(SLATE)
val.major_gridlines.format.line.color.rgb = rgb("E3E7EC")
val.format.line.fill.background()
scaling = val._element.find(qn("c:scaling"))
log_base = scaling.makeelement(qn("c:logBase"), {"val": "10"})
scaling.insert(0, log_base)  # schema: logBase comes first in c:scaling
val.minimum_scale = 1
val.has_title = True
val.axis_title.text_frame.text = "mm"
val.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)

# per-sample slides
for x in ok:
    s = prs.slides.add_slide(BLANK)
    background(s, PAPER)
    text(s, 0.5, 0.3, 6, 0.7, x["subject"], size=36, colour=INK, font=HEAD, bold=True)
    text(s, 0.5, 0.97, 8.3, 0.3, x["asset"], size=11, colour=MUTED)
    main = x["images"].get("map") or x["images"].get("overview")
    if main:
        w = 8.3
        h = w * main["h"] / main["w"]
        if h > 5.85:
            h, w = 5.85, 5.85 * main["w"] / main["h"]
        s.shapes.add_picture(main["path"], Inches(0.5), Inches(1.35), Inches(w), Inches(h))
    cards = (
        (f"{fmt(x['axon_mm'], 1 if (x['axon_mm'] or 0) < 10 else 0)} mm",
         f"axon candidates · {x['axon_pieces']} SWCs", AXON),
        (f"{fmt(x['dendrite_mm'])} mm", "dendrites of detected somata", DENDRITE_TEXT),
        (f"{fmt(x['skeleton_mm'], 1 if x['skeleton_mm'] < 10 else 0)} mm",
         f"skeleton · {x['pieces']} pieces", INK),
        (fmt(x["somata"]), "somata near the skeleton", INK),
    )  # fmt: skip
    for i, (value, label, colour) in enumerate(cards):
        cx, cy = 9.1 + (i % 2) * 1.9, 1.35 + (i // 2) * 1.2
        card(s, cx, cy, 1.75, 1.05)
        text(s, cx + 0.12, cy + 0.1, 1.55, 0.5, value, size=22, colour=colour, font=HEAD, bold=True)
        text(s, cx + 0.12, cy + 0.6, 1.55, 0.4, label, size=10, colour=SLATE)
    if x["images"].get("map") and x["images"].get("overview"):
        o = x["images"]["overview"]
        h = min(3.65 * o["h"] / o["w"], 3.2)
        w = h * o["w"] / o["h"]
        s.shapes.add_picture(o["path"], Inches(9.1), Inches(3.85), Inches(w), Inches(h))
        text(s, 9.1, 3.9 + h, 3.65, 0.3, "On the fused image: region outline, pieces, close-ups",
             size=10, colour=MUTED)  # fmt: skip
    s.notes_slide.notes_text_frame.text = (
        f"{x['subject']}: {fmt(x['skeleton_mm'], 1)} mm of skeleton in the grown MD "
        f"({fmt(x['region_mm3'], 1)} mm³); {fmt(x['dendrite_mm'], 1)} mm dendrite of "
        f"{x['somata'] or 0} detected somata; {fmt(x['axon_mm'], 1)} mm axon candidates in "
        f"{x['axon_pieces']} SWCs."
    )

# closing
s = prs.slides.add_slide(BLANK)
background(s, INK)
text(s, 0.8, 0.8, 11.7, 0.8, "Not in this deck", size=36, colour=PAPER, font=HEAD, bold=True)
notes = [
    f"• {', '.join(f['subject'] for f in failed) or 'none'}: run failed (its CCF annotation is "
    "missing a pyramid level; it also has no soma detections)",
    "• 39 more exaSPIM subjects: no processed asset with both a fused image and a CCF annotation",
    "• Samples with little cable (e.g. 773889, 773890) look sparsely labelled in MD; their faint "
    "fibres need a closer look",
]  # fmt: skip
box = text(s, 0.8, 1.9, 11.7, 3.5, notes, size=18, colour=LIGHT)
for p in box.text_frame.paragraphs:
    p.space_after = Pt(12)

prs.save(sys.argv[2])
print("wrote", sys.argv[2], len(prs.slides), "slides")
